import csv
import io
from datetime import timedelta

import jdatetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from ..db import get_db
from ..deps import audit, require
from ..models import Alarm, User
from ..notifier import KIND_FA
from .alarms import filtered

from fastapi import Request

router = APIRouter(prefix="/api/reports", tags=["reports"])
STATUS_FA = {"new": "جدید", "ack": "رسیدگی شد", "false_positive": "هشدار کاذب"}
SEV_FA = {"high": "بالا", "medium": "متوسط"}
HEAD = ["شناسه", "تاریخ (شمسی)", "ساعت", "دوربین", "نوع تخلف", "شدت", "مدت (ثانیه)", "وضعیت", "رسیدگی‌کننده", "توضیحات"]
FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def _rows(db, frm, to, camera_id, status, tz_offset):
    off = timedelta(minutes=-tz_offset)
    out = []
    for a in filtered(db, status, None, camera_id, frm, to).order_by(Alarm.id.desc()).limit(50000).all():
        loc = a.created_at + off
        j = jdatetime.datetime.fromgregorian(datetime=loc)
        out.append([a.id, j.strftime("%Y/%m/%d"), loc.strftime("%H:%M:%S"), a.camera.name, KIND_FA.get(a.kind, a.kind),
                    SEV_FA.get(a.severity, a.severity), round(a.duration, 1), STATUS_FA.get(a.status, a.status),
                    a.handled_by, a.note])
    return out


@router.get("/alarms.{fmt}")
def export(fmt: str, request: Request, frm: str | None = None, to: str | None = None, camera_id: int | None = None,
           status: str | None = None, tz_offset: int = -210, user: User = Depends(require("reports.export")),
           db: Session = Depends(get_db)):
    if fmt not in ("csv", "xlsx"):
        raise HTTPException(404, "فرمت پشتیبانی نمی‌شود")
    rows = _rows(db, frm, to, camera_id, status, tz_offset)
    audit(db, user, "report_export", f"{fmt} rows={len(rows)}", request)
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(HEAD)
        w.writerows(rows)
        return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": 'attachment; filename="safetycheck-alarms.csv"'})
    wb = Workbook()
    ws = wb.active
    ws.title = "آلارم‌ها"
    ws.sheet_view.rightToLeft = True
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2A44")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for col, width in zip("ABCDEFGHIJ", [8, 14, 11, 22, 24, 10, 12, 14, 16, 36]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    bio = io.BytesIO()
    wb.save(bio)
    return Response(bio.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="safetycheck-alarms.xlsx"'})
